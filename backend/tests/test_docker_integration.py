"""Run with DOCMIND_INTEGRATION=1 against docker compose services."""
import base64
import os
import subprocess
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import httpx
import jwt
import pytest
from openpyxl import load_workbook
from pypdf import PdfWriter

BASE_URL = os.getenv("DOCMIND_BASE_URL", "http://localhost:8000")
pytestmark = pytest.mark.skipif(os.getenv("DOCMIND_INTEGRATION") != "1", reason="Docker integration only")


@pytest.fixture(autouse=True)
def reset_redis() -> None:
    """Integration tests own Redis DB 0; clear rate windows and queued test tasks."""
    subprocess.run(
        ["docker", "compose", "exec", "-T", "redis", "redis-cli", "FLUSHDB"],
        check=True,
        capture_output=True,
        cwd=Path(__file__).parents[2],
    )


def request(method: str, path: str, **kwargs: object) -> httpx.Response:
    return httpx.request(method, f"{BASE_URL}{path}", timeout=30, **kwargs)


def account(label: str) -> tuple[str, str]:
    email = f"{label.lower().replace(' ', '-')}-{uuid.uuid4().hex[:8]}@example.com"
    response = request("POST", "/api/v1/auth/register", json={"email": email, "password": "correct-horse-battery-staple", "organization_name": label})
    assert response.status_code == 201, response.text
    data = response.json()
    return data["access_token"], data["organization_id"]


def auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def pdf(pages: int = 1) -> bytes:
    writer = PdfWriter()
    writer.add_metadata({"/DocMindTestId": uuid.uuid4().hex})
    for _ in range(pages):
        writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def project(token: str, organization_id: str) -> str:
    response = request("POST", f"/api/v1/projects?organization_id={organization_id}&name=integration", headers=auth(token))
    assert response.status_code == 201, response.text
    return response.json()["id"]


def upload(token: str, organization_id: str, project_id: str, pages: int = 1) -> dict:
    response = request(
        "POST",
        f"/api/v1/documents?organization_id={organization_id}&project_id={project_id}",
        headers=auth(token),
        files={"file": ("fixture.pdf", pdf(pages), "application/pdf")},
    )
    assert response.status_code == 202, response.text
    return response.json()


def wait_complete(token: str, organization_id: str, document_id: str) -> None:
    for _ in range(20):
        response = request("GET", f"/api/v1/documents/{document_id}/status?organization_id={organization_id}", headers=auth(token))
        assert response.status_code == 200, response.text
        if response.json()["status"] == "COMPLETED":
            return
        time.sleep(1)
    pytest.fail("worker did not complete document")


def postgres_scalar(sql: str) -> str:
    return subprocess.check_output(
        [
            "docker",
            "compose",
            "exec",
            "-T",
            "db",
            "psql",
            "-U",
            "docmind",
            "-d",
            "docmind",
            "-At",
            "-c",
            sql,
        ],
        text=True,
        cwd=Path(__file__).parents[2],
    ).strip()


def make_storage_failure(storage_key: str) -> None:
    subprocess.run(
        [
            "docker", "compose", "exec", "-T", "api", "python", "-c",
            "from app.storage import storage; import sys; p=storage._path(sys.argv[1]); p.unlink(); p.mkdir()",
            storage_key,
        ],
        check=True,
        capture_output=True,
        cwd=Path(__file__).parents[2],
    )


def restore_storage(storage_key: str, content: bytes) -> None:
    subprocess.run(
        [
            "docker", "compose", "exec", "-T", "api", "python", "-c",
            "from app.storage import storage; import base64,sys; p=storage._path(sys.argv[1]); p.rmdir(); p.write_bytes(base64.b64decode(sys.argv[2]))",
            storage_key,
            base64.b64encode(content).decode(),
        ],
        check=True,
        capture_output=True,
        cwd=Path(__file__).parents[2],
    )


def wait_for_job_state(job_id: str, prefix: str) -> str:
    for _ in range(250):
        state = postgres_scalar(
            "SELECT status || '|' || COALESCE(failure_code, '') "
            f"FROM processing_jobs WHERE id = '{job_id}'"
        )
        if state.startswith(prefix):
            return state
        time.sleep(0.1)
    pytest.fail(f"job {job_id} did not reach {prefix}")


def test_jwt_idor_api_key_scope_and_xlsx_export() -> None:
    token_a, org_a = account("Org A")
    token_b, org_b = account("Org B")
    project_a = project(token_a, org_a)
    document = upload(token_a, org_a, project_a)
    wait_complete(token_a, org_a, document["document_id"])

    cross_tenant = request("GET", f"/api/v1/documents/{document['document_id']}?organization_id={org_b}", headers=auth(token_b))
    assert cross_tenant.status_code in {403, 404}
    unauthorized_export = request("GET", f"/api/v1/documents/{document['document_id']}/export?organization_id={org_b}&format=xlsx", headers=auth(token_b))
    assert unauthorized_export.status_code in {403, 404}

    key_response = request("POST", f"/api/v1/api-keys?organization_id={org_a}&name=reader", headers=auth(token_a), json=["documents:read"])
    assert key_response.status_code == 201, key_response.text
    key = key_response.json()["key"]
    assert request("GET", f"/api/v1/documents?organization_id={org_a}", headers=auth(key)).status_code == 200
    assert request("GET", f"/api/v1/usage?organization_id={org_a}", headers=auth(key)).status_code == 403
    assert request("GET", f"/api/v1/documents?organization_id={org_b}", headers=auth(key)).status_code == 403

    export_key = request("POST", f"/api/v1/api-keys?organization_id={org_a}&name=exporter", headers=auth(token_a), json=["documents:read", "exports:read"])
    assert export_key.status_code == 201
    xlsx = request("GET", f"/api/v1/documents/{document['document_id']}/export?organization_id={org_a}&format=xlsx", headers=auth(export_key.json()["key"]))
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    workbook = load_workbook(BytesIO(xlsx.content))
    assert workbook.active["A1"].value == "Field"


def test_free_quota_and_concurrent_reservation() -> None:
    token, organization_id = account("Quota")
    project_id = project(token, organization_id)
    assert upload(token, organization_id, project_id, 10)["status"] == "QUEUED"
    assert upload(token, organization_id, project_id, 10)["status"] == "QUEUED"
    blocked = request(
        "POST",
        f"/api/v1/documents?organization_id={organization_id}&project_id={project_id}",
        headers=auth(token),
        files={"file": ("one-more.pdf", pdf(1), "application/pdf")},
    )
    assert blocked.status_code == 402

    race_token, race_org = account("Quota Race")
    race_project = project(race_token, race_org)

    def reserve() -> int:
        return request(
            "POST",
            f"/api/v1/documents?organization_id={race_org}&project_id={race_project}",
            headers=auth(race_token),
            files={"file": (f"race-{uuid.uuid4().hex}.pdf", pdf(12), "application/pdf")},
        ).status_code

    with ThreadPoolExecutor(max_workers=2) as executor:
        outcomes = list(executor.map(lambda _: reserve(), range(2)))
    assert sorted(outcomes) == [202, 402]


def test_login_rate_limit() -> None:
    responses = [
        request("POST", "/api/v1/auth/login", json={"email": "nobody@example.com", "password": "not-a-real-password"})
        for _ in range(11)
    ]
    assert [response.status_code for response in responses[:-1]] == [401] * 10
    assert responses[-1].status_code == 429
    assert responses[-1].headers["retry-after"] == "60"


def test_invalid_credentials_and_upload_validation() -> None:
    token, organization_id = account("Upload Security")
    project_id = project(token, organization_id)
    valid = request("GET", f"/api/v1/documents?organization_id={organization_id}", headers=auth(token))
    assert valid.status_code == 200
    for bad_token in ["not-a-jwt", "abc.def.ghi", "dm_test_missing"]:
        assert request("GET", f"/api/v1/documents?organization_id={organization_id}", headers=auth(bad_token)).status_code == 401

    malformed = request(
        "POST",
        f"/api/v1/documents?organization_id={organization_id}&project_id={project_id}",
        headers=auth(token),
        files={"file": ("not-a-pdf.pdf", b"not a PDF", "application/pdf")},
    )
    assert malformed.status_code == 415

    oversized = request(
        "POST",
        f"/api/v1/documents?organization_id={organization_id}&project_id={project_id}",
        headers=auth(token),
        files={"file": ("large.pdf", b"%PDF-" + b"0" * (26_214_401), "application/pdf")},
    )
    assert oversized.status_code == 413

    names = [
        "../../evil.pdf", "../../../etc/passwd", "..\\..\\evil.pdf", "....//....//evil.pdf",
        "%2e%2e%2fevil.pdf", "/tmp/evil.pdf", "C:\\temp\\evil.pdf", "<script>alert(1)</script>.pdf",
        '"; DROP TABLE documents; --.pdf', "  unicode-ñ-文書.pdf  ", f"{'x' * 300}.pdf",
        "document.pdf.exe", "document.exe.pdf",
    ]
    storage_keys: list[str] = []
    for name in names:
        response = request(
            "POST",
            f"/api/v1/documents?organization_id={organization_id}&project_id={project_id}",
            headers=auth(token),
            files={"file": (name, pdf(), "application/pdf")},
        )
        expected = 202 if name.lower().endswith(".pdf") else 415
        assert response.status_code == expected, response.text
        if expected != 202:
            continue
        key = postgres_scalar(
            "SELECT storage_key FROM documents "
            f"WHERE id = '{response.json()['document_id']}'"
        )
        assert key.startswith(f"{organization_id}/")
        assert ".." not in key and not key.startswith(("/", "\\"))
        storage_keys.append(key)
    stored = subprocess.run(
        [
            "docker", "compose", "exec", "-T", "api", "python", "-c",
            "from app.storage import storage; import sys; assert all(storage._path(k).is_file() for k in sys.argv[1:])",
            *storage_keys,
        ],
        capture_output=True,
        check=False,
        cwd=Path(__file__).parents[2],
    )
    assert stored.returncode == 0


def test_jwt_and_api_key_expiration() -> None:
    token, organization_id = account("Expiration")
    assert request("GET", f"/api/v1/documents?organization_id={organization_id}", headers=auth(token)).status_code == 200
    payload = jwt.decode(token, options={"verify_signature": False})
    expired_token = subprocess.check_output(
        [
            "docker", "compose", "exec", "-T", "api", "python", "-c",
            "import jwt,sys; from datetime import UTC,datetime,timedelta; from app.core import settings; print(jwt.encode({'sub':sys.argv[1],'exp':datetime.now(UTC)-timedelta(seconds=1)},settings().jwt_secret,algorithm=settings().jwt_algorithm))",
            payload["sub"],
        ],
        text=True,
        cwd=Path(__file__).parents[2],
    ).strip()
    assert request("GET", f"/api/v1/documents?organization_id={organization_id}", headers=auth(expired_token)).status_code == 401
    wrong_signature = jwt.encode({"sub": payload["sub"], "exp": int(time.time()) + 60}, "wrong-test-secret-with-at-least-thirty-two-bytes", algorithm="HS256")
    assert request("GET", f"/api/v1/documents?organization_id={organization_id}", headers=auth(wrong_signature)).status_code == 401

    created = request("POST", f"/api/v1/api-keys?organization_id={organization_id}&name=expiry", headers=auth(token), json=["documents:read"])
    assert created.status_code == 201
    key = created.json()
    assert request("GET", f"/api/v1/documents?organization_id={organization_id}", headers=auth(key["key"])).status_code == 200
    subprocess.run(
        ["docker", "compose", "exec", "-T", "db", "psql", "-U", "docmind", "-d", "docmind", "-c", f"UPDATE api_keys SET expires_at = NOW() - INTERVAL '1 second' WHERE id = '{key['id']}';"],
        check=True,
        capture_output=True,
        cwd=Path(__file__).parents[2],
    )
    assert request("GET", f"/api/v1/documents?organization_id={organization_id}", headers=auth(key["key"])).status_code == 401

    revocable = request(
        "POST",
        f"/api/v1/api-keys?organization_id={organization_id}&name=revocable",
        headers=auth(token),
        json=["documents:read"],
    )
    assert revocable.status_code == 201
    revocable_key = revocable.json()
    assert request(
        "GET",
        f"/api/v1/documents?organization_id={organization_id}",
        headers=auth(revocable_key["key"]),
    ).status_code == 200
    assert request(
        "DELETE",
        f"/api/v1/api-keys/{revocable_key['id']}?organization_id={organization_id}",
        headers=auth(token),
    ).status_code == 204
    assert request(
        "GET",
        f"/api/v1/documents?organization_id={organization_id}",
        headers=auth(revocable_key["key"]),
    ).status_code == 401


def test_worker_failure_is_safe_and_does_not_charge_usage() -> None:
    token, organization_id = account("Worker Failure")
    project_id = project(token, organization_id)
    document_id, job_id = str(uuid.uuid4()), str(uuid.uuid4())
    missing_key = f"{organization_id}/missing-{uuid.uuid4()}.pdf"
    postgres_scalar(
        "INSERT INTO documents "
        "(id, organization_id, project_id, filename, storage_key, mime_type, checksum, created_at) "
        f"VALUES ('{document_id}', '{organization_id}', '{project_id}', "
        f"'failure-fixture.pdf', '{missing_key}', 'application/pdf', '{uuid.uuid4().hex}', NOW()); "
        "INSERT INTO processing_jobs (id, organization_id, document_id, status, attempt, created_at) "
        f"VALUES ('{job_id}', '{organization_id}', '{document_id}', 'queued', 1, NOW());"
    )
    subprocess.run(
        [
            "docker",
            "compose",
            "exec",
            "-T",
            "api",
            "python",
            "-c",
            "from app.worker import process_document; process_document.delay(__import__('sys').argv[1])",
            job_id,
        ],
        check=True,
        capture_output=True,
        cwd=Path(__file__).parents[2],
    )
    for _ in range(20):
        state = postgres_scalar(
            "SELECT status || '|' || COALESCE(failure_message, '') || '|' || attempt "
            f"FROM processing_jobs WHERE id = '{job_id}'"
        )
        if state.startswith("failed|"):
            break
        time.sleep(1)
    assert state == "failed|Processing failed. Check operational logs.|1"
    assert postgres_scalar(
        "SELECT status || '|' || COALESCE(error, '') "
        f"FROM processing_steps WHERE job_id = '{job_id}' AND name = 'text_extraction'"
    ) == "FAILED|PROCESSING_FAILED"
    assert postgres_scalar(
        "SELECT count(*) FROM usage_records "
        f"WHERE document_id = '{document_id}' AND metric IN ('documents_processed', 'pages_processed')"
    ) == "0"
    assert postgres_scalar(
        f"SELECT count(*) FROM extraction_runs WHERE document_id = '{document_id}'"
    ) == "0"


def test_process_endpoint_is_idempotent_under_concurrency() -> None:
    token, organization_id = account("Process Idempotency")
    project_id = project(token, organization_id)
    document = upload(token, organization_id, project_id)
    wait_complete(token, organization_id, document["document_id"])

    def process() -> httpx.Response:
        return request(
            "POST",
            f"/api/v1/documents/{document['document_id']}/process?organization_id={organization_id}",
            headers=auth(token),
        )

    with ThreadPoolExecutor(max_workers=10) as executor:
        responses = list(executor.map(lambda _: process(), range(10)))
    assert all(response.status_code == 202 for response in responses)
    job_ids = {response.json()["job_id"] for response in responses}
    assert len(job_ids) == 1
    assert sum(not response.json()["reused"] for response in responses) == 1
    wait_complete(token, organization_id, document["document_id"])
    assert postgres_scalar(
        "SELECT count(*) FROM processing_jobs "
        f"WHERE document_id = '{document['document_id']}'"
    ) == "2"
    assert postgres_scalar(
        "SELECT count(*) FROM extraction_runs "
        f"WHERE document_id = '{document['document_id']}'"
    ) == "2"


def test_api_key_upload_process_extraction_export_and_revoke() -> None:
    token, organization_id = account("API Key E2E")
    project_id = project(token, organization_id)
    created = request(
        "POST",
        f"/api/v1/api-keys?organization_id={organization_id}&name=e2e",
        headers=auth(token),
        json=["documents:read", "documents:write", "exports:read"],
    )
    assert created.status_code == 201
    key = created.json()
    document = upload(key["key"], organization_id, project_id)
    wait_complete(key["key"], organization_id, document["document_id"])
    processed = request(
        "POST",
        f"/api/v1/documents/{document['document_id']}/process?organization_id={organization_id}",
        headers=auth(key["key"]),
    )
    assert processed.status_code == 202
    wait_complete(key["key"], organization_id, document["document_id"])
    extraction = request(
        "GET",
        f"/api/v1/documents/{document['document_id']}/extraction?organization_id={organization_id}",
        headers=auth(key["key"]),
    )
    assert extraction.status_code == 200
    xlsx = request(
        "GET",
        f"/api/v1/documents/{document['document_id']}/export?organization_id={organization_id}&format=xlsx",
        headers=auth(key["key"]),
    )
    assert xlsx.status_code == 200
    assert load_workbook(BytesIO(xlsx.content)).active["A1"].value == "Field"
    assert request(
        "DELETE",
        f"/api/v1/api-keys/{key['id']}?organization_id={organization_id}",
        headers=auth(token),
    ).status_code == 204
    assert request(
        "GET",
        f"/api/v1/documents?organization_id={organization_id}",
        headers=auth(key["key"]),
    ).status_code == 401


def test_transient_storage_failure_retries_once_without_duplicate_usage() -> None:
    token, organization_id = account("Retry Success")
    project_id = project(token, organization_id)
    document = upload(token, organization_id, project_id)
    wait_complete(token, organization_id, document["document_id"])
    storage_key = postgres_scalar(
        f"SELECT storage_key FROM documents WHERE id = '{document['document_id']}'"
    )
    replacement = pdf()
    make_storage_failure(storage_key)
    queued = request(
        "POST",
        f"/api/v1/documents/{document['document_id']}/process?organization_id={organization_id}",
        headers=auth(token),
    )
    assert queued.status_code == 202
    job_id = queued.json()["job_id"]
    assert wait_for_job_state(job_id, "queued|PROCESSING_RETRY") == "queued|PROCESSING_RETRY"
    restore_storage(storage_key, replacement)
    wait_complete(token, organization_id, document["document_id"])
    failed_attempts = int(postgres_scalar(
        "SELECT count(*) FROM processing_steps "
        f"WHERE job_id = '{job_id}' AND name = 'text_extraction' AND status = 'FAILED'"
    ))
    assert 1 <= failed_attempts <= 2
    assert postgres_scalar(
        "SELECT count(*) FROM extraction_runs "
        f"WHERE document_id = '{document['document_id']}'"
    ) == "2"
    assert postgres_scalar(
        "SELECT COALESCE(sum(quantity), 0) FROM usage_records "
        f"WHERE document_id = '{document['document_id']}' AND metric = 'pages_processed'"
    ) == "2"


def test_transient_failure_stops_after_max_retries() -> None:
    token, organization_id = account("Retry Maximum")
    project_id = project(token, organization_id)
    document = upload(token, organization_id, project_id)
    wait_complete(token, organization_id, document["document_id"])
    storage_key = postgres_scalar(
        f"SELECT storage_key FROM documents WHERE id = '{document['document_id']}'"
    )
    make_storage_failure(storage_key)
    queued = request(
        "POST",
        f"/api/v1/documents/{document['document_id']}/process?organization_id={organization_id}",
        headers=auth(token),
    )
    assert queued.status_code == 202
    job_id = queued.json()["job_id"]
    assert wait_for_job_state(job_id, "failed|PROCESSING_FAILED") == "failed|PROCESSING_FAILED"
    assert postgres_scalar(
        "SELECT count(*) FROM processing_steps "
        f"WHERE job_id = '{job_id}' AND name = 'text_extraction' AND status = 'FAILED'"
    ) == "4"
    assert postgres_scalar(
        "SELECT count(*) FROM extraction_runs "
        f"WHERE document_id = '{document['document_id']}'"
    ) == "1"
    assert postgres_scalar(
        "SELECT COALESCE(sum(quantity), 0) FROM usage_records "
        f"WHERE document_id = '{document['document_id']}' AND metric = 'pages_processed'"
    ) == "1"
    restore_storage(storage_key, pdf())


def test_retry_reprocess_reuses_the_waiting_job() -> None:
    token, organization_id = account("Retry Reprocess")
    project_id = project(token, organization_id)
    document = upload(token, organization_id, project_id)
    wait_complete(token, organization_id, document["document_id"])
    storage_key = postgres_scalar(
        f"SELECT storage_key FROM documents WHERE id = '{document['document_id']}'"
    )
    make_storage_failure(storage_key)
    first = request(
        "POST",
        f"/api/v1/documents/{document['document_id']}/process?organization_id={organization_id}",
        headers=auth(token),
    )
    assert first.status_code == 202
    assert wait_for_job_state(first.json()["job_id"], "queued|PROCESSING_RETRY")
    duplicate = request(
        "POST",
        f"/api/v1/documents/{document['document_id']}/process?organization_id={organization_id}",
        headers=auth(token),
    )
    assert duplicate.status_code == 202
    assert duplicate.json()["reused"] is True
    assert duplicate.json()["job_id"] == first.json()["job_id"]
    restore_storage(storage_key, pdf())
    wait_complete(token, organization_id, document["document_id"])
    assert postgres_scalar(
        "SELECT count(*) FROM processing_jobs "
        f"WHERE document_id = '{document['document_id']}'"
    ) == "2"
    assert postgres_scalar(
        "SELECT count(*) FROM extraction_runs "
        f"WHERE document_id = '{document['document_id']}'"
    ) == "2"
