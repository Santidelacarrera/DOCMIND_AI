"""End-to-end API behaviour: auth, tenancy, uploads, exports and hardening."""

import csv
import io
import json
import uuid

import jwt
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app import main
from app.db import SessionLocal
from app.models import ExtractionField, ExtractionRun, RunStatus
from tests.conftest import PASSWORD, Account, make_pdf


# ---------------------------------------------------------------- authentication
def test_register_login_me_logout_invalidates_tokens(client: TestClient, register) -> None:
    acct = register()
    login = client.post("/api/v1/auth/login", json={"email": acct.email, "password": PASSWORD})
    assert login.status_code == 200
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    assert client.get("/api/v1/auth/me", headers=headers).json()["email"] == acct.email
    assert client.post("/api/v1/auth/logout", headers=headers).status_code == 204
    client.cookies.clear()
    assert client.get("/api/v1/auth/me", headers=headers).status_code == 401


def test_wrong_and_unknown_credentials_are_indistinguishable(client: TestClient, register) -> None:
    acct = register()
    bad_pw = client.post("/api/v1/auth/login", json={"email": acct.email, "password": "x" * 14})
    unknown = client.post(
        "/api/v1/auth/login", json={"email": "nobody@example.com", "password": "x" * 14}
    )
    assert bad_pw.status_code == unknown.status_code == 401
    assert bad_pw.json() == unknown.json()


def test_weak_password_and_duplicate_email_rejected(client: TestClient, register) -> None:
    acct = register()
    short = client.post(
        "/api/v1/auth/register",
        json={"email": "a@example.com", "password": "short", "organization_name": "x"},
    )
    assert short.status_code == 422
    dup = client.post(
        "/api/v1/auth/register",
        json={"email": acct.email.upper(), "password": PASSWORD, "organization_name": "x"},
    )
    assert dup.status_code == 409


def test_jwt_edge_cases_never_500(client: TestClient) -> None:
    secret = "test-secret-test-secret-test-secret-123"
    cases = {
        "garbage": "not-a-token",
        "alg-none": jwt.encode({"sub": str(uuid.uuid4()), "sv": 1}, key="", algorithm="none"),
        "bad-sub": jwt.encode({"sub": "not-a-uuid", "sv": 1, "exp": 9999999999}, secret, "HS256"),
        "no-exp": jwt.encode({"sub": str(uuid.uuid4()), "sv": 1}, secret, "HS256"),
        "unknown-user": jwt.encode(
            {"sub": str(uuid.uuid4()), "sv": 1, "exp": 9999999999}, secret, "HS256"
        ),
        "wrong-secret": jwt.encode(
            {"sub": str(uuid.uuid4()), "sv": 1, "exp": 9999999999}, "x" * 40, "HS256"
        ),
    }
    for name, token in cases.items():
        response = client.get("/api/v1/auth/me", headers={"Authorization": f"Bearer {token}"})
        assert response.status_code == 401, name


def test_login_rate_limit_per_account_and_fail_closed(
    client: TestClient, register, fake_redis, monkeypatch: pytest.MonkeyPatch
) -> None:
    acct = register()
    monkeypatch.setenv("RATE_LIMIT_LOGIN", "3")
    from app.core import settings

    settings.cache_clear()
    try:
        codes = [
            client.post(
                "/api/v1/auth/login", json={"email": acct.email, "password": "w" * 14}
            ).status_code
            for _ in range(5)
        ]
        assert codes[:3] == [401, 401, 401]
        assert codes[3:] == [429, 429]
    finally:
        monkeypatch.delenv("RATE_LIMIT_LOGIN")
        settings.cache_clear()
    fake_redis.unavailable = True
    assert (
        client.post(
            "/api/v1/auth/login", json={"email": acct.email, "password": PASSWORD}
        ).status_code
        == 503
    )


def test_cookie_session_requires_csrf(client: TestClient, register) -> None:
    acct = register()
    login = client.post("/api/v1/auth/login", json={"email": acct.email, "password": PASSWORD})
    assert login.status_code == 200
    url = f"/api/v1/projects?organization_id={acct.org}&name=p"
    assert client.post(url).status_code == 403  # cookie present, no CSRF header
    csrf = client.cookies.get("csrf_token")
    assert client.post(url, headers={"X-CSRF-Token": csrf}).status_code == 201
    assert client.post(url, headers={"X-CSRF-Token": "wrong"}).status_code == 403
    # Logout is also protected against cross-site requests.
    assert client.post("/api/v1/auth/logout").status_code == 403


def test_security_headers_and_no_docs_exposure_flags(client: TestClient) -> None:
    response = client.get("/health")
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "DENY"
    assert "frame-ancestors 'none'" in response.headers["content-security-policy"]


def test_cors_is_restricted_to_configured_origin(client: TestClient) -> None:
    allowed = client.options(
        "/api/v1/auth/login",
        headers={"Origin": "http://localhost:3000", "Access-Control-Request-Method": "POST"},
    )
    assert allowed.headers.get("access-control-allow-origin") == "http://localhost:3000"
    denied = client.options(
        "/api/v1/auth/login",
        headers={"Origin": "https://evil.example", "Access-Control-Request-Method": "POST"},
    )
    assert "access-control-allow-origin" not in denied.headers


# ---------------------------------------------------------------- tenancy
def test_tenant_isolation_across_every_resource(client: TestClient, register) -> None:
    alice, mallory = register("alice"), register("mallory")
    project = alice.project()
    doc = alice.upload(project).json()["document_id"]
    spy = {"Authorization": f"Bearer {mallory.token}"}
    # Mallory addresses Alice's org directly and via her own org with Alice's ids.
    for org in (alice.org, mallory.org):
        for path in (
            f"/api/v1/documents?organization_id={org}",
            f"/api/v1/documents/{doc}?organization_id={org}",
            f"/api/v1/documents/{doc}/download?organization_id={org}",
            f"/api/v1/documents/{doc}/status?organization_id={org}",
            f"/api/v1/documents/{doc}/extraction?organization_id={org}",
            f"/api/v1/projects?organization_id={org}",
            f"/api/v1/audit-logs?organization_id={org}",
            f"/api/v1/usage?organization_id={org}",
            f"/api/v1/api-keys?organization_id={org}",
        ):
            response = client.get(path, headers=spy)
            if org == alice.org:
                assert response.status_code == 403, path
            else:
                assert doc not in response.text, path
    stolen = client.post(
        f"/api/v1/documents/{doc}/process?organization_id={mallory.org}", headers=spy
    )
    assert stolen.status_code == 404
    upload_to_foreign_project = mallory.upload(project)
    assert upload_to_foreign_project.status_code == 404


def test_invalid_organization_id_is_rejected(client: TestClient, register) -> None:
    acct = register()
    assert (
        client.get("/api/v1/projects?organization_id=nope", headers=acct.headers).status_code == 400
    )


def test_viewer_cannot_write(client: TestClient, register) -> None:
    from app.models import OrganizationMember, Role, User

    owner, viewer = register("owner"), register("viewer")
    with SessionLocal.begin() as db:
        user = db.scalar(select(User).where(User.email == viewer.email))
        db.add(OrganizationMember(organization_id=uuid.UUID(owner.org), user_id=user.id, role=Role.viewer))
    response = client.post(
        f"/api/v1/projects?organization_id={owner.org}&name=x", headers=viewer.headers
    )
    assert response.status_code == 403
    assert (
        client.get(f"/api/v1/projects?organization_id={owner.org}", headers=viewer.headers).status_code
        == 200
    )


# ---------------------------------------------------------------- API keys
def test_api_key_lifecycle_and_scope_enforcement(client: TestClient, register) -> None:
    acct = register()
    project = acct.project()
    created = client.post(
        f"/api/v1/api-keys?organization_id={acct.org}&name=ci",
        json=["documents:read"],
        headers=acct.headers,
    )
    assert created.status_code == 201
    secret, key_id = created.json()["key"], created.json()["id"]
    assert secret.startswith("dm_live_")
    key = {"Authorization": f"Bearer {secret}"}
    assert client.get(f"/api/v1/documents?organization_id={acct.org}", headers=key).status_code == 200
    # Scope missing -> forbidden; user-only endpoints -> forbidden, never 500.
    assert acct.upload(project).status_code == 202
    denied = client.post(
        f"/api/v1/documents?organization_id={acct.org}&project_id={project}",
        files={"file": ("a.pdf", make_pdf(), "application/pdf")},
        headers=key,
    )
    assert denied.status_code == 403
    for method, path in (
        ("get", "/api/v1/auth/me"),
        ("get", "/api/v1/organizations"),
        ("post", "/api/v1/auth/logout"),
        ("post", "/api/v1/organizations?name=x"),
        ("get", f"/api/v1/api-keys?organization_id={acct.org}"),
        ("get", f"/api/v1/audit-logs?organization_id={acct.org}"),
        ("post", f"/api/v1/api-keys?organization_id={acct.org}&name=n"),
    ):
        extra = {"json": ["documents:read"]} if "api-keys" in path and method == "post" else {}
        assert getattr(client, method)(path, headers=key, **extra).status_code == 403, path
    assert (
        client.delete(
            f"/api/v1/api-keys/{key_id}?organization_id={acct.org}", headers=acct.headers
        ).status_code
        == 204
    )
    assert client.get(f"/api/v1/documents?organization_id={acct.org}", headers=key).status_code == 401


def test_api_key_cannot_cross_tenants_and_rejects_unknown_scopes(client: TestClient, register) -> None:
    a, b = register("a"), register("b")
    bad = client.post(
        f"/api/v1/api-keys?organization_id={a.org}&name=n",
        json=["admin:everything"],
        headers=a.headers,
    )
    assert bad.status_code == 422
    secret = client.post(
        f"/api/v1/api-keys?organization_id={a.org}&name=n",
        json=["documents:read"],
        headers=a.headers,
    ).json()["key"]
    other = client.get(
        f"/api/v1/documents?organization_id={b.org}", headers={"Authorization": f"Bearer {secret}"}
    )
    assert other.status_code == 403


# ---------------------------------------------------------------- uploads
def test_upload_validations(client: TestClient, register) -> None:
    acct = register()
    project = acct.project()
    assert acct.upload(project, b"not a pdf", "x.pdf").status_code == 415
    assert acct.upload(project, make_pdf(), "x.txt").status_code == 415
    assert acct.upload(project, b"%PDF-1.4 truncated garbage", "x.pdf").status_code == 422
    good = make_pdf()
    first = acct.upload(project, good)
    assert first.status_code == 202
    assert acct.upload(project, good).status_code == 409
    assert len(main.queued_jobs) == 1  # type: ignore[attr-defined]


def test_upload_enforces_size_pages_and_quota(
    client: TestClient, register, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.core import settings

    acct = register()
    project = acct.project()
    monkeypatch.setattr(settings(), "max_upload_bytes", 200)
    assert acct.upload(project, make_pdf()).status_code == 413
    monkeypatch.setattr(settings(), "max_upload_bytes", 26_214_400)
    monkeypatch.setattr(settings(), "max_pdf_pages", 2)
    assert acct.upload(project, make_pdf(3)).status_code == 422
    monkeypatch.setattr(settings(), "max_pdf_pages", 100)
    monkeypatch.setattr(settings(), "free_pages_per_month", 3)
    assert acct.upload(project, make_pdf(2)).status_code == 202
    assert acct.upload(project, make_pdf(2)).status_code == 402


def test_malicious_filenames_are_sanitized(client: TestClient, register) -> None:
    acct = register()
    project = acct.project()
    response = acct.upload(project, make_pdf(), '../../etc/pa"ss\r\nX-Evil: 1.pdf')
    assert response.status_code == 202
    doc_id = response.json()["document_id"]
    listing = client.get(f"/api/v1/documents?organization_id={acct.org}", headers=acct.headers)
    name = listing.json()[0]["filename"]
    assert "/" not in name and "\r" not in name and "\n" not in name
    download = client.get(
        f"/api/v1/documents/{doc_id}/download?organization_id={acct.org}", headers=acct.headers
    )
    assert download.status_code == 200
    assert "x-evil" not in download.headers
    assert "\n" not in download.headers["content-disposition"]
    assert download.headers["content-disposition"].startswith("attachment;")
    inline = client.get(
        f"/api/v1/documents/{doc_id}/download?organization_id={acct.org}&inline=true",
        headers=acct.headers,
    )
    assert inline.headers["content-disposition"].startswith("inline;")
    assert "frame-ancestors http://localhost:3000" in inline.headers["content-security-policy"]


def test_upload_populates_page_count_and_process_is_idempotent(client: TestClient, register) -> None:
    acct = register()
    project = acct.project()
    body = acct.upload(project, make_pdf(3)).json()
    listing = client.get(f"/api/v1/documents?organization_id={acct.org}", headers=acct.headers)
    assert listing.json()[0]["pages"] == 3
    again = client.post(
        f"/api/v1/documents/{body['document_id']}/process?organization_id={acct.org}",
        headers=acct.headers,
    )
    assert again.json()["reused"] is True and again.json()["job_id"] == body["job_id"]


# ---------------------------------------------------------------- extraction + export
def _seed_extraction(acct: Account, document_id: str, fields: dict[str, object]) -> list[str]:
    with SessionLocal.begin() as db:
        run = ExtractionRun(
            organization_id=uuid.UUID(acct.org),
            document_id=uuid.UUID(document_id),
            provider="mock",
            status=RunStatus.completed,
            result={"fields": fields},
        )
        db.add(run)
        db.flush()
        ids = []
        for name, value in fields.items():
            field = ExtractionField(extraction_run_id=run.id, name=name, original_value=value, value=value)
            db.add(field)
            db.flush()
            ids.append(str(field.id))
        return ids


def test_exports_neutralize_spreadsheet_formulas(client: TestClient, register) -> None:
    acct = register()
    doc = acct.upload(acct.project()).json()["document_id"]
    _seed_extraction(
        acct,
        doc,
        {"=evil": "=HYPERLINK(\"http://x\")", "plain": "ok", "nested": {"a": 1}, "num": 5},
    )
    base = f"/api/v1/documents/{doc}/export?organization_id={acct.org}"
    rows = list(csv.DictReader(io.StringIO(client.get(base + "&format=csv", headers=acct.headers).text)))
    by_name = {r["name"]: r["value"] for r in rows}
    assert by_name["'=evil"].startswith("'=HYPERLINK")
    assert by_name["plain"] == "ok" and json.loads(by_name["nested"]) == {"a": 1}
    workbook = load_workbook(io.BytesIO(client.get(base + "&format=xlsx", headers=acct.headers).content))
    cells = [c.value for row in workbook.active.iter_rows(min_row=2) for c in row]
    assert not any(isinstance(v, str) and v.startswith("=") for v in cells)
    as_json = client.get(base, headers=acct.headers).json()
    assert {f["name"] for f in as_json["fields"]} == {"=evil", "plain", "nested", "num"}
    assert client.get(base + "&format=exe", headers=acct.headers).status_code == 422


def test_edit_field_marks_verified_and_is_tenant_scoped(client: TestClient, register) -> None:
    a, b = register("a"), register("b")
    doc = a.upload(a.project()).json()["document_id"]
    (field_id,) = _seed_extraction(a, doc, {"total": 10})
    url = f"/api/v1/extraction-fields/{field_id}?organization_id="
    assert client.patch(url + b.org, json=99, headers=b.headers).status_code == 404
    assert client.patch(url + a.org, json=99, headers=b.headers).status_code == 403
    ok = client.patch(url + a.org, json={"x": [1, 2]}, headers=a.headers)
    assert ok.status_code == 200 and ok.json()["manually_verified"] is True
    huge = client.patch(url + a.org, json="x" * 70_000, headers=a.headers)
    assert huge.status_code == 413


# ---------------------------------------------------------------- unit helpers
@pytest.mark.parametrize("raw", [None, "", "..", "../../x.pdf", "a\\b\\c.pdf", "\x00\x01.pdf", "é" * 400])
def test_sanitize_filename_is_always_safe(raw: str | None) -> None:
    name = main.sanitize_filename(raw)
    assert name and len(name) <= 255
    assert "/" not in name and "\\" not in name and name.isprintable()


@pytest.mark.parametrize("value", ["=1+1", "+1", "-1", "@SUM(A1)", "\tx", "\rx"])
def test_neutralize_formula_prefixes(value: str) -> None:
    assert main.neutralize_formula(value) == "'" + value
    assert main.neutralize_formula("safe") == "safe"
    assert main.neutralize_formula(5) == 5


def test_health_and_ready(client: TestClient) -> None:
    assert client.get("/health").json() == {"status": "ok"}
    assert client.get("/ready").json() == {"status": "ready"}
